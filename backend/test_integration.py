import sys
import os
import urllib.request
import json
import uuid

# Set absolute path for backend import (to create dummy docx)
sys.path.append(r'c:\Users\Admin\Desktop\SkillGapAI\backend')

try:
    import docx
except ImportError:
    print("python-docx not installed.")
    sys.exit(1)

print("Starting full-stack integration test verification...")

# 1. Create a dummy docx file
doc_path = "integration_test_resume.docx"
doc = docx.Document()
doc.add_heading('Jane Doe Resume', level=0)
doc.add_paragraph('Skills: Python, Git, HTML5, CSS3, communication, problem solving')
doc.add_paragraph('Experience: Developer at WebStudio')
doc.save(doc_path)

try:
    # Read binary bytes of the docx file
    with open(doc_path, 'rb') as f:
        docx_bytes = f.read()

    # 2. Upload file to running Flask server
    url = "http://127.0.0.1:5000/api/resume/upload"
    boundary = b'Boundary-' + uuid.uuid4().hex.encode('utf-8')
    
    body = (
        b'--' + boundary + b'\r\n' +
        b'Content-Disposition: form-data; name="file"; filename="integration_test_resume.docx"\r\n' +
        b'Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n' +
        docx_bytes + b'\r\n' +
        b'--' + boundary + b'--\r\n'
    )
    
    headers = {
        'Content-Type': f'multipart/form-data; boundary={boundary.decode("utf-8")}',
        'Content-Length': str(len(body))
    }
    
    req = urllib.request.Request(url, data=body, headers=headers, method='POST')
    print("Uploading DOCX resume to http://127.0.0.1:5000/api/resume/upload ...")
    with urllib.request.urlopen(req) as response:
        res_data = json.loads(response.read().decode('utf-8'))
        print("Upload Response:", json.dumps(res_data, indent=2))
        assert res_data.get('status') == 'success', "Upload failed"
        extracted_skills = res_data.get('data', {}).get('extracted_skills', [])
        print(f"Extracted Skills Count: {len(extracted_skills)}")

    # 3. Choose a Career Goal and Generate Roadmap
    url_goal = "http://127.0.0.1:5000/api/roadmap/generate"
    goal_payload = json.dumps({"goal_title": "Backend Engineer"}).encode('utf-8')
    req_goal = urllib.request.Request(
        url_goal, 
        data=goal_payload, 
        headers={'Content-Type': 'application/json'}, 
        method='POST'
    )
    
    print("\nGenerating roadmap for target role: Backend Engineer ...")
    with urllib.request.urlopen(req_goal) as response:
        res_goal = json.loads(response.read().decode('utf-8'))
        print("Roadmap Response Match Skills:", res_goal.get('data', {}).get('skills_matched', []))
        print("Roadmap Response Missing Skills Gaps:", res_goal.get('data', {}).get('skills_missing', []))
        assert res_goal.get('status') == 'success', "Roadmap generation failed"
        
    # 4. Get active statistics dashboard
    url_stats = "http://127.0.0.1:5000/api/progress/stats"
    req_stats = urllib.request.Request(url_stats, method='GET')
    
    print("\nQuerying progress stats dashboard ...")
    with urllib.request.urlopen(req_stats) as response:
        res_stats = json.loads(response.read().decode('utf-8'))
        data = res_stats.get('data', {})
        print(f"Dashboard readiness score: {data.get('readiness_score')}%")
        print(f"Roadmap progress completed nodes: {data.get('roadmap_progress')}%")
        assert res_stats.get('status') == 'success', "Stats lookup failed"

    print("\nINTEGRATION TEST SUCCESSFULLY VERIFIED!")

except Exception as e:
    print(f"\nERROR: Integration test failed: {str(e)}")
    sys.exit(1)

finally:
    # Clean up
    if os.path.exists(doc_path):
        os.remove(doc_path)
